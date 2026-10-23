#!/usr/bin/env python3
"""
Merged-heading sensitivity (Option 3, steps 4-6). Applies the 11 approved MERGE pairs from the reviewed crosswalk.
Safeguard: headings merge into one harmonised identity; their NDC episodes consolidate into one spell only when the
headings' observation windows overlap. Non-overlapping windows stay separate spells under the same identity.
The shared exit rule is reapplied across the consolidated NDC population. Everything else is identical to
compare_reconstructions.py, so the two outputs are directly comparable.
"""
import csv, re, sys, os, collections, statistics as st, datetime as dt
import openpyxl
sys.path.insert(0, ".")
from stage1_join_audit import parse_date

APPROVED = {  # reviewed 23 Oct 2026; 11 pairs
    ("Bumetanide Injection", "Bumetanide Injection, USP"), ("Bupivacaine Hydrochloride Injection", "Bupivacaine Hydrochloride Injection, USP"),
    ("Capecitabine Tablet", "Capecitabine Tablets"), ("Furosemide Injection", "Furosemide Injection, USP"),
    ("Hydrocortisone Tablets", "Hydrocortisone Tablets, USP"), ("Iomeprol Injection", "Iomeprol injection"),
    ("Midazolam Injection", "Midazolam Injection, USP"), ("Nitroglycerin Injectable", "Nitroglycerin Injection"),
    ("Semaglutide (Ozempic) Injection", "Semaglutide Injection"), ("0.9% Sodium Chloride Irrigation", "Sodium Chloride 0.9% Irrigation"),
    ("Sulfasalazine Tablet", "Sulfasalazine Tablets"),
}
X = list(csv.DictReader(open("stage4_compare/heading_dedup_crosswalk.csv")))
pairs = [(r["heading_a"], r["heading_b"]) for r in X if (r["heading_a"], r["heading_b"]) in APPROVED or (r["heading_b"], r["heading_a"]) in APPROVED]
assert len(pairs) == 11, pairs

OUT = "stage4_compare/sensitivity_merged"; os.makedirs(OUT, exist_ok=True)
wb = openpyxl.load_workbook("FDA_Recovery_Time_Validation.xlsx", read_only=True)
rows = list(wb["FDA Product Episodes"].iter_rows(values_only=True)); H = rows[3]
EP = [dict(zip(H, r)) for r in rows[4:] if r and r[0]]
for e in EP:
    for k in ("Initial posting", "Resolution / censor", "First seen", "Last seen"): e[k] = parse_date(e[k])
snap = sorted({parse_date(r[1]) for r in list(wb["Snapshot Audit"].iter_rows(values_only=True))[4:] if r and r[1] and parse_date(r[1])})
FIRST, LAST = snap[0], snap[-1]
def prev_snap(d):
    p = [s for s in snap if s < d]; return max(p) if p else None

# ---------------------------------------------------------------- harmonised identity (union-find over approved pairs)
parent = {}
def find(x):
    parent.setdefault(x, x)
    while parent[x] != x: parent[x] = parent[parent[x]]; x = parent[x]
    return x
for a, b in pairs: parent[find(a)] = find(b)
ident = {h: find(h) for h in {e["Generic name"] for e in EP}}
label = {}
for h, r in ident.items(): label.setdefault(r, min([k for k, v in ident.items() if v == r], key=len))

# ---------------------------------------------------------------- spells: consolidate overlapping windows within an identity
by_head = collections.defaultdict(list)
for e in EP: by_head[e["Generic name"]].append(e)
spells = []; merge_log = []
for r in set(ident.values()):
    heads = [h for h, v in ident.items() if v == r]
    if len(heads) == 1:
        spells.append((label[r], heads, by_head[heads[0]])); continue
    # cluster headings whose observation windows overlap
    win = {h: (min(x["First seen"] for x in by_head[h]), max(x["Last seen"] for x in by_head[h])) for h in heads}
    clusters = []
    for h in sorted(heads, key=lambda h: win[h][0]):
        for c in clusters:
            if any(win[h][0] <= win[g][1] and win[g][0] <= win[h][1] for g in c): c.append(h); break
        else: clusters.append([h])
    for c in clusters:
        eps = [e for h in c for e in by_head[h]]
        spells.append((label[r] + ("" if len(clusters) == 1 else f" [spell {clusters.index(c)+1}]"), c, eps))
        merge_log.append({"harmonised_identity": label[r], "headings_consolidated": " | ".join(c), "n_spells_for_identity": len(clusters), "ndc_rows": len(eps),
                          "action": "consolidated: windows overlap" if len(c) > 1 else "kept separate spell: window does not overlap other heading(s)"})

drugs = []
for name, heads, v in spells:
    posting = min(x["Initial posting"] for x in v); fs = min(x["First seen"] for x in v); ls = max(x["Last seen"] for x in v)
    resolved = [x for x in v if x["Event observed"]]
    if len(resolved) == len(v): exit_status, exit_date = "RESOLVED", max(x["Resolution / censor"] for x in v)
    elif resolved: exit_status, exit_date = "AMBIGUOUS", None
    elif ls >= LAST: exit_status, exit_date = "ONGOING", LAST
    else: exit_status, exit_date = "DISAPPEARED", ls
    drugs.append(dict(generic=name, headings=len(heads), n_ndc=len(v), posting=posting, first_seen=fs, last_seen=ls, prev=prev_snap(fs) if fs > FIRST else None,
                      exit_status=exit_status, exit_date=exit_date, event=int(exit_status == "RESOLVED"), r2_eligible=posting >= FIRST and fs > FIRST))

# ---------------------------------------------------------------- identical machinery to compare_reconstructions.py
def km(durs, events):
    pairs_ = sorted(zip(durs, events)); S = 1.0; curve = []; n = len(pairs_); i = 0
    while i < n:
        t = pairs_[i][0]; tied = [x for x in pairs_[i:] if x[0] == t]; d = sum(x[1] for x in tied)
        if d: S *= 1 - d / (n - i)
        curve.append((t, S)); i += len(tied)
    return curve
def S_at(c, t):
    s = 1.0
    for x, v in c:
        if x <= t: s = v
        else: break
    return s
def km_median(c): return next((t for t, s in c if s < 0.5), "not reached")
elig = [d for d in drugs if d["r2_eligible"]]; amb = [d for d in elig if d["exit_status"] == "AMBIGUOUS"]; common = [d for d in elig if d["exit_status"] != "AMBIGUOUS"]
def dur(d, onset): return (d["exit_date"] - onset).days
def latest_onset(d):
    c = [d["first_seen"]]
    if d["exit_status"] == "RESOLVED": c.append(d["exit_date"])
    return min(c)
R1 = [dur(d, d["posting"]) for d in common]; EV = [d["event"] for d in common]
R2e = [dur(d, d["prev"]) for d in common]; R2l = [dur(d, latest_onset(d)) for d in common]
def summ(name, D):
    c = km(D, EV); return dict(rule=name, spells=len(D), events=sum(EV), KM_median=km_median(c), S110=round(S_at(c, 110), 3), S180=round(S_at(c, 180), 3), S365=round(S_at(c, 365), 3), gt1y=round(sum(1 for x in D if x > 365) / len(D), 3))
T = [summ("R1", R1), summ("R2 earliest", R2e), summ("R2 latest", R2l)]
def top20(D): return {common[i]["generic"] for i in sorted(range(len(common)), key=lambda i: -D[i])[:20]}
T1, T2e, T2l = top20(R1), top20(R2e), top20(R2l); U = T1 | T2e | T2l; DEF = T1 & T2e & T2l
def spearman(x, y):
    def rk(v):
        o = sorted(range(len(v)), key=lambda i: v[i]); r = [0] * len(v)
        for pos, i in enumerate(o): r[i] = pos
        return r
    rx, ry = rk(x), rk(y); mx, my = st.mean(rx), st.mean(ry)
    return sum((a - mx) * (b - my) for a, b in zip(rx, ry)) / ((sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry)) ** 0.5)
diff = [a - b for a, b in zip(R1, R2l)]

# ---------------------------------------------------------------- attribution of the original four ranking changes
orig = list(csv.DictReader(open("stage4_compare/rank_stability.csv")))
orig_changes = [r["generic"] for r in orig if r["tier"] == "POSSIBLE"]
merged_heads = {h for a, b in pairs for h in (a, b)}
attribution = [{"original_POSSIBLE_drug": g, "involves_duplicate_heading": int(g in merged_heads), "harmonised_identity": label.get(ident.get(g, g), g),
                "tier_after_merge": "DEFINITE" if label.get(ident.get(g, g), g) in DEF else "POSSIBLE" if label.get(ident.get(g, g), g) in U else "not top-20"} for g in orig_changes]

# ---------------------------------------------------------------- write
with open(f"{OUT}/merge_log.csv", "w", newline="") as f: w = csv.DictWriter(f, fieldnames=list(merge_log[0].keys())); w.writeheader(); w.writerows(merge_log)
with open(f"{OUT}/paired_durations_merged.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["identity", "headings", "exit_status", "event", "R1_days", "R2_earliest_days", "R2_latest_days"])
    for i, d in enumerate(common): w.writerow([d["generic"], d["headings"], d["exit_status"], d["event"], R1[i], R2e[i], R2l[i]])
with open(f"{OUT}/rank_stability_merged.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["identity", "tier", "in_R1", "in_R2e", "in_R2l", "R1_days", "R2_latest_days"])
    for i, d in enumerate(common):
        if d["generic"] in U: w.writerow([d["generic"], "DEFINITE" if d["generic"] in DEF else "POSSIBLE", int(d["generic"] in T1), int(d["generic"] in T2e), int(d["generic"] in T2l), R1[i], R2l[i]])
with open(f"{OUT}/attribution_of_original_changes.csv", "w", newline="") as f: w = csv.DictWriter(f, fieldnames=list(attribution[0].keys())); w.writeheader(); w.writerows(attribution)

orig_T = {r["rule"]: r for r in csv.DictReader(open("stage4_compare/table3_reconstruction_summaries.csv"))}
md = f"""# Merged-heading sensitivity (11 approved duplicate pairs; temporal-spell safeguard applied)

Approved pairs: {len(pairs)}; headings involved: {len(merged_heads)}; consolidation log in merge_log.csv
({sum(1 for m in merge_log if m['action'].startswith('consolidated'))} identities consolidated into one spell, {sum(1 for m in merge_log if m['action'].startswith('kept'))} kept as separate spells).

| Result | Original headings | Merged-heading sensitivity |
|---|---|---|
| Administrative units (headings / harmonised spells) | 298 | {len(drugs)} |
| Entry-identifiable cohort | 184 | {len(elig)} |
| Analytical observations | 173 | {len(common)} |
| Resolved events | 51 | {sum(EV)} |
| Ambiguous exits (in cohort) | 11 | {len(amb)} |
| Paired R1 − R2(latest): median, IQR | 37, 16–118 | {round(st.median(diff))}, {round(st.quantiles(diff, n=4)[0])}–{round(st.quantiles(diff, n=4)[2])} |
| Drugs with R1 − R2(latest) > 365 d | 11 | {sum(1 for x in diff if x > 365)} |
| One-year survival, R1 / R2e / R2l | 0.823 / 0.840 / 0.789 | {T[0]['S365']} / {T[1]['S365']} / {T[2]['S365']} |
| KM median, R1 / R2e / R2l | NR / NR / 916 | {T[0]['KM_median']} / {T[1]['KM_median']} / {T[2]['KM_median']} |
| Stable top-20 membership (DEFINITE) | 16/20 | {len(DEF)}/20 |
| Top-20 candidate union (DEFINITE + POSSIBLE) | 24 | {len(U)} ({len(DEF)} + {len(U - DEF)}) |
| Spearman R1 vs R2-latest / R2-earliest | 0.825 / 0.817 | {round(spearman(R1, R2l), 3)} / {round(spearman(R1, R2e), 3)} |

## Which identities leave / enter the top 20 under archive-observed onset (merged units)
- Leave (in R1 only): {sorted(T1 - T2e - T2l)}
- Enter (in R2 only): {sorted((T2e & T2l) - T1)}

## Attribution of the original four ranking changes
{chr(10).join(f"- {a['original_POSSIBLE_drug'][:40]}: duplicate heading involved = {a['involves_duplicate_heading']}; after merge → {a['tier_after_merge']}" for a in attribution)}
"""
open(f"{OUT}/sensitivity_summary.md", "w").write(md); print(md)
